from __future__ import annotations

from io import BytesIO
import zipfile


from miy_api.domains.document_processing import extract_document
from miy_api.domains.document_processing.extractors import (
    DocumentExtractBundle,
    EvidenceBlock,
)


def _pptx_bytes(entries: dict[str, str]) -> bytes:
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for path, xml in entries.items():
            archive.writestr(path, xml)
    return buffer.getvalue()


def _xlsx_bytes(rows: list[list[str]]) -> bytes:
    from openpyxl import Workbook

    buffer = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    workbook.save(buffer)
    return buffer.getvalue()


def _docx_with_malformed_vertical_merge_row() -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    document.add_paragraph("고장과 무관한 본문 절차")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "병합 셀"
    table.cell(0, 1).text = "보존할 표 값"
    table.cell(1, 0).text = "병합 연속"
    table.cell(1, 1).text = "마지막 값"
    table.cell(0, 0).merge(table.cell(1, 0))

    # Reproduce a real-world malformed vertical merge: the continuation cell
    # points to grid offset 0, but the row above declares that offset as absent.
    first_row = table._tbl.tr_lst[0]  # pyright: ignore[reportPrivateUsage]
    first_row.remove(first_row.tc_lst[0])
    grid_before = OxmlElement("w:gridBefore")
    grid_before.set(qn("w:val"), "1")
    first_row.get_or_add_trPr().append(grid_before)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _slide_xml(*paragraphs: str) -> str:
    body = "".join(f"<a:p><a:r><a:t>{paragraph}</a:t></a:r></a:p>" for paragraph in paragraphs)
    return (
        '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f"<p:cSld><p:spTree>{body}</p:spTree></p:cSld></p:sld>"
    )


def _notes_xml(*paragraphs: str) -> str:
    return _slide_xml(*paragraphs)


def _slide_with_table_xml() -> str:
    return (
        '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        "<p:cSld><p:spTree>"
        "<a:p><a:r><a:t>제품 사양</a:t></a:r></a:p>"
        "<a:tbl>"
        "<a:tr>"
        "<a:tc><a:txBody><a:p><a:r><a:t>항목</a:t></a:r></a:p></a:txBody></a:tc>"
        "<a:tc><a:txBody><a:p><a:r><a:t>값</a:t></a:r></a:p></a:txBody></a:tc>"
        "</a:tr>"
        "<a:tr>"
        "<a:tc><a:txBody><a:p><a:r><a:t>정격 출력</a:t></a:r></a:p></a:txBody></a:tc>"
        "<a:tc><a:txBody><a:p><a:r><a:t>10 kW</a:t></a:r></a:p></a:txBody></a:tc>"
        "</a:tr>"
        "</a:tbl>"
        "</p:spTree></p:cSld></p:sld>"
    )


class FixedExtractor:
    def supports(self, *, mime_type: str, filename: str) -> bool:
        return filename == "custom.fixture"

    def extract(
        self,
        *,
        document_id: str,
        filename: str,
        mime_type: str,
        content: bytes,
    ) -> DocumentExtractBundle:
        return DocumentExtractBundle(
            document_id=document_id,
            filename=filename,
            mime_type=mime_type,
            evidence_blocks=[
                EvidenceBlock(
                    document_id=document_id,
                    block_id=f"{document_id}:custom",
                    locator_kind="custom",
                    locator_label="Custom",
                    section_path="Custom",
                    block_kind="text",
                    text=content.decode("utf-8"),
                )
            ],
        )


def test_extract_embedded_office_documents_skips_decompression_bomb() -> None:
    from miy_api.domains.document_processing.extractors import (
        extract_embedded_office_documents,
    )

    # A tiny compressed member that expands ~1000x. It must be rejected by the
    # size/expansion-ratio guard before archive.read() allocates it in full,
    # never decompressed into a result.
    outer = BytesIO()
    with zipfile.ZipFile(outer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("ppt/presentation.xml", "<presentation />")
        archive.writestr("ppt/embeddings/bomb.pptx", b"\x00" * (5 * 1024 * 1024))

    docs = extract_embedded_office_documents(outer.getvalue())

    assert docs == []


def test_html_extractor_preserves_structured_locators_and_strips_active_content() -> None:
    content = """
    <!doctype html>
    <html>
      <head>
        <title>제동 시스템 시험 결과</title>
        <style>.secret { display: block }</style>
        <script src="https://example.invalid/execute.js">실행 금지 스크립트</script>
      </head>
      <body>
        <h1>요약</h1>
        <p>정상 제동 압력은 12 bar 입니다.</p>
        <ul><li><p>시험 장비 교정 완료</p></li></ul>
        <table>
          <tr><th>항목</th><th>결과</th></tr>
          <tr><td>응답 시간</td><td>120 ms</td></tr>
        </table>
        <iframe src="https://example.invalid/remote">원격 프레임 본문</iframe>
        <div hidden>숨김 본문</div>
        <noscript>대체 실행 본문</noscript>
      </body>
    </html>
    """.encode()

    bundle = extract_document(
        document_id="html-doc",
        filename="report.html",
        mime_type="text/html",
        content=content,
    )

    assert [block.locator_kind for block in bundle.evidence_blocks] == [
        "title",
        "heading",
        "paragraph",
        "list_item",
        "table_row",
        "table_row",
    ]
    assert bundle.evidence_blocks[2].section_path == "요약"
    assert bundle.evidence_blocks[4].rows == [["항목", "결과"]]
    assert bundle.evidence_blocks[5].locator_label == "Table 1, row 2"
    assert "실행 금지" not in bundle.normalized_text
    assert "원격 프레임" not in bundle.normalized_text
    assert "숨김 본문" not in bundle.normalized_text
    assert "대체 실행" not in bundle.normalized_text
    assert "example.invalid" not in bundle.normalized_text
    assert bundle.metadata["format"] == "html"
    assert bundle.metadata["charset"] == "utf-8-sig"
    assert bundle.metadata["table_row_count"] == 2
    assert bundle.metadata["truncated"] is False
